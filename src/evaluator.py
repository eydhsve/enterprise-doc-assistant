"""Offline retrieval evaluator for CSV Q&A datasets."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from .config import Settings
from .vector_store import ChromaVectorStore, VectorStoreError


logger = logging.getLogger(__name__)


class EvaluationError(RuntimeError):
    """Raised when an evaluation dataset or run is invalid."""


@dataclass
class EvaluationResult:
    """Summary and per-question evaluation tables."""

    summary: pd.DataFrame
    details: pd.DataFrame
    output_path: str | None = None

    def summary_text(self) -> str:
        if self.summary.empty:
            return "没有可展示的评测结果。"
        row = self.summary.iloc[0]
        return (
            f"样本数: {int(row['样本数'])} | "
            f"命中数: {int(row['命中数'])} | "
            f"Hit@{int(row['K'])}: {float(row['Hit@K']):.4f} | "
            f"MRR: {float(row['MRR']):.4f}"
        )


class RetrievalEvaluator:
    """Compute Top-K retrieval hit rate without invoking an LLM."""

    REQUIRED_COLUMNS = {"question", "expected_source"}

    def __init__(
        self,
        settings: Settings,
        vector_store: ChromaVectorStore,
    ) -> None:
        self.settings = settings
        self.vector_store = vector_store

    def evaluate(
        self,
        dataset_path: str | Path,
        *,
        top_k: int = 5,
        output_path: str | Path | None = None,
    ) -> EvaluationResult:
        """Read CSV, run retrieval, and optionally export XLSX tables."""

        dataset = self._load_dataset(dataset_path)
        requested_k = max(1, int(top_k))
        detail_rows: list[dict[str, Any]] = []

        for row_index, row in dataset.iterrows():
            question = str(row.get("question", "")).strip()
            expected_source = str(row.get("expected_source", "")).strip()
            expected_answer = str(row.get("expected_answer", "")).strip()

            if not question:
                detail_rows.append(
                    self._error_row(
                        row_index,
                        question,
                        expected_source,
                        expected_answer,
                        requested_k,
                        "question 为空。",
                    )
                )
                continue

            try:
                hits = self.vector_store.search(question, top_k=requested_k)
            except VectorStoreError as exc:
                detail_rows.append(
                    self._error_row(
                        row_index,
                        question,
                        expected_source,
                        expected_answer,
                        requested_k,
                        str(exc),
                    )
                )
                continue

            retrieved_sources = [
                str(hit.metadata.get("source", "")) for hit in hits
            ]
            retrieved_filenames = [
                str(hit.metadata.get("filename", "")) for hit in hits
            ]
            hit_rank = self._first_hit_rank(
                expected_source,
                retrieved_sources,
                retrieved_filenames,
            )
            hit = hit_rank is not None

            detail_rows.append(
                {
                    "序号": int(row_index) + 1,
                    "问题": question,
                    "期望来源": expected_source,
                    "期望答案": expected_answer,
                    "K": requested_k,
                    "是否命中": bool(hit),
                    "首次命中排名": hit_rank,
                    "倒数排名": 1.0 / hit_rank if hit_rank else 0.0,
                    "检索来源": " | ".join(retrieved_sources),
                    "检索文件": " | ".join(retrieved_filenames),
                    "错误": "",
                }
            )

        details = pd.DataFrame(detail_rows)
        if details.empty:
            details = pd.DataFrame(
                columns=[
                    "序号",
                    "问题",
                    "期望来源",
                    "期望答案",
                    "K",
                    "是否命中",
                    "首次命中排名",
                    "倒数排名",
                    "检索来源",
                    "检索文件",
                    "错误",
                ]
            )

        valid_rows = details[details["错误"].fillna("").astype(str) == ""]
        total = int(len(valid_rows))
        hit_count = int(valid_rows["是否命中"].sum()) if total else 0
        hit_rate = hit_count / total if total else 0.0
        mrr = float(valid_rows["倒数排名"].mean()) if total else 0.0

        summary = pd.DataFrame(
            [
                {
                    "数据集": str(Path(dataset_path)),
                    "K": requested_k,
                    "样本数": total,
                    "命中数": hit_count,
                    "错误数": int(len(details) - total),
                    "Hit@K": hit_rate,
                    "MRR": mrr,
                    "评测时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
            ]
        )

        result = EvaluationResult(summary=summary, details=details)
        if output_path is not None:
            saved = self.save_results(result, output_path)
            result.output_path = str(saved)
        return result

    def save_results(
        self,
        result: EvaluationResult,
        output_path: str | Path,
    ) -> Path:
        """Write professional XLSX summary and detail sheets."""

        path = Path(output_path).expanduser()
        if not path.is_absolute():
            path = self.settings.eval_output_dir / path
        if path.suffix.lower() != ".xlsx":
            path = path.with_suffix(".xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)

        try:
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                result.summary.to_excel(
                    writer,
                    sheet_name="评测指标",
                    index=False,
                )
                result.details.to_excel(
                    writer,
                    sheet_name="问题明细",
                    index=False,
                )
            self._format_workbook(path)
        except Exception as exc:
            raise EvaluationError(f"评测表格导出失败: {exc}") from exc
        return path.resolve()

    def default_output_path(self) -> Path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return self.settings.eval_output_dir / f"retrieval_eval_{timestamp}.xlsx"

    @staticmethod
    def _load_dataset(path: str | Path) -> pd.DataFrame:
        dataset_path = Path(path).expanduser()
        try:
            dataset_path = dataset_path.resolve(strict=True)
        except FileNotFoundError as exc:
            raise EvaluationError(f"评测数据集不存在: {dataset_path}") from exc
        if not dataset_path.is_file():
            raise EvaluationError(f"评测数据集不是文件: {dataset_path}")
        if dataset_path.suffix.lower() != ".csv":
            raise EvaluationError("评测数据集必须是 CSV 文件。")

        try:
            dataset = pd.read_csv(dataset_path, encoding="utf-8-sig")
        except UnicodeDecodeError:
            try:
                dataset = pd.read_csv(dataset_path, encoding="gb18030")
            except Exception as exc:
                raise EvaluationError(
                    f"CSV 编码无法识别: {dataset_path}"
                ) from exc
        except pd.errors.EmptyDataError as exc:
            raise EvaluationError("评测 CSV 为空。") from exc
        except pd.errors.ParserError as exc:
            raise EvaluationError(f"CSV 格式解析失败: {exc}") from exc
        except OSError as exc:
            raise EvaluationError(f"CSV 读取失败: {exc}") from exc

        dataset.columns = [str(column).strip() for column in dataset.columns]
        missing = RetrievalEvaluator.REQUIRED_COLUMNS - set(dataset.columns)
        if missing:
            raise EvaluationError(
                "CSV 缺少必要列: " + ", ".join(sorted(missing))
            )
        if dataset.empty:
            raise EvaluationError("评测 CSV 没有数据行。")

        if "expected_answer" not in dataset.columns:
            dataset["expected_answer"] = ""
        return dataset

    @staticmethod
    def _first_hit_rank(
        expected_source: str,
        retrieved_sources: list[str],
        retrieved_filenames: list[str],
    ) -> int | None:
        expected_items = [
            item.strip().lower()
            for item in re.split(r"[|;]+", expected_source or "")
            if item.strip()
        ]
        if not expected_items:
            return None

        for rank, (source, filename) in enumerate(
            zip(retrieved_sources, retrieved_filenames),
            start=1,
        ):
            for expected in expected_items:
                if RetrievalEvaluator.source_matches(
                    expected,
                    source,
                    filename,
                ):
                    return rank
        return None

    @staticmethod
    def source_matches(
        expected: str,
        retrieved_source: str,
        retrieved_filename: str = "",
    ) -> bool:
        expected_value = (expected or "").strip().lower()
        source_value = (retrieved_source or "").strip().lower()
        filename_value = (retrieved_filename or "").strip().lower()
        if not expected_value:
            return False
        if not source_value and not filename_value:
            return False
        if expected_value in {source_value, filename_value}:
            return True
        if not filename_value and source_value:
            filename_value = Path(source_value).name.lower()
        checks = []
        if source_value:
            checks.extend(
                [
                    expected_value in source_value,
                    source_value in expected_value,
                ]
            )
        if filename_value:
            checks.extend(
                [
                    expected_value in filename_value,
                    filename_value in expected_value,
                ]
            )
        return any(checks)

    @staticmethod
    def _error_row(
        row_index: int,
        question: str,
        expected_source: str,
        expected_answer: str,
        top_k: int,
        error: str,
    ) -> dict[str, Any]:
        return {
            "序号": int(row_index) + 1,
            "问题": question,
            "期望来源": expected_source,
            "期望答案": expected_answer,
            "K": top_k,
            "是否命中": False,
            "首次命中排名": None,
            "倒数排名": 0.0,
            "检索来源": "",
            "检索文件": "",
            "错误": error,
        }

    @staticmethod
    def _format_workbook(path: Path) -> None:
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            raise EvaluationError(
                "缺少 openpyxl，无法格式化评测表格。"
            ) from exc

        workbook = load_workbook(path)
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        header_font = Font(name="Arial", bold=True, color="1F2937")
        body_font = Font(name="Arial", color="111827")

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = body_font
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

            for column_cells in worksheet.columns:
                max_length = max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in column_cells
                )
                column_letter = column_cells[0].column_letter
                worksheet.column_dimensions[column_letter].width = min(
                    max(max_length + 2, 12),
                    50,
                )

        summary_sheet = workbook["评测指标"]
        for cell in summary_sheet[2]:
            if cell.column_letter in {"F", "G"}:
                cell.number_format = "0.00%"

        workbook.save(path)
