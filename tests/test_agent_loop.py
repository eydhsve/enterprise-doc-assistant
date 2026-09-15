"""Function-calling loop tests using local fakes."""

from __future__ import annotations

import json
import tempfile
import unittest
from collections import deque
from pathlib import Path
from typing import Any

from src.agent_tools import DocumentAgent
from src.config import Settings
from src.evaluator import RetrievalEvaluator
from src.vector_store import SearchHit


def make_settings(root: Path) -> Settings:
    data_dir = root / "data"
    return Settings(
        project_root=root,
        zhipu_api_key="test-key",
        zhipu_base_url="https://open.bigmodel.cn/api/paas/v4",
        zhipu_model="glm-4.7-flash",
        llm_temperature=0.2,
        llm_timeout=30.0,
        llm_max_retries=0,
        embedding_model_name="BAAI/bge-small-zh",
        embedding_device="cpu",
        embedding_local_only=True,
        embedding_cache_dir=data_dir / "models",
        chroma_persist_dir=data_dir / "chroma",
        chroma_collection="test_collection",
        chunk_size=500,
        chunk_overlap=100,
        retrieval_top_k=3,
        max_history_messages=6,
        max_context_chars=8000,
        upload_dir=data_dir / "uploads",
        report_dir=data_dir / "reports",
        eval_output_dir=data_dir / "evaluations",
        gradio_server_name="127.0.0.1",
        gradio_server_port=7860,
        gradio_share=False,
    )


class FakeVectorStore:
    def search(
        self,
        query: str,
        *,
        top_k: int | None = None,
        source_filter: str | None = None,
    ) -> list[SearchHit]:
        return [
            SearchHit(
                vector_id="doc-1:0",
                text="员工通过试用期后可以申请年假。",
                metadata={
                    "doc_id": "doc-1",
                    "source": "/docs/employee_handbook.md",
                    "filename": "employee_handbook.md",
                    "chunk_index": 0,
                },
                score=0.93,
                distance=0.07,
            )
        ][: top_k or 3]


class FakeLLM:
    def __init__(self) -> None:
        self.responses = deque(
            [
                {
                    "role": "assistant",
                    "content": "",
                    "tool_calls": [
                        {
                            "id": "call_retrieve",
                            "type": "function",
                            "function": {
                                "name": "knowledge_retrieve",
                                "arguments": json.dumps(
                                    {"query": "员工何时申请年假"},
                                    ensure_ascii=False,
                                ),
                            },
                        }
                    ],
                },
                {
                    "role": "assistant",
                    "content": "",
                    "tool_calls": [
                        {
                            "id": "call_export",
                            "type": "function",
                            "function": {
                                "name": "export_report",
                                "arguments": json.dumps(
                                    {
                                        "question": "员工何时申请年假？",
                                        "answer": "通过试用期后可以申请。",
                                        "sources": ["employee_handbook.md"],
                                        "report_name": "annual_leave",
                                    },
                                    ensure_ascii=False,
                                ),
                            },
                        }
                    ],
                },
                {
                    "role": "assistant",
                    "content": "员工通过试用期后可以申请年假。",
                    "tool_calls": [],
                },
            ]
        )
        self.calls: list[dict[str, Any]] = []

    def chat_completion(
        self,
        messages: list[dict[str, Any]],
        **kwargs: Any,
    ) -> dict[str, Any]:
        self.calls.append({"messages": list(messages), "kwargs": dict(kwargs)})
        return self.responses.popleft()


class AgentLoopTests(unittest.TestCase):
    def test_agent_executes_tools_and_exports_report(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            settings = make_settings(root)
            settings.ensure_directories()
            fake_llm = FakeLLM()
            agent = DocumentAgent(
                settings,
                llm_client=fake_llm,  # type: ignore[arg-type]
                vector_store=FakeVectorStore(),  # type: ignore[arg-type]
            )

            response = agent.run("员工何时申请年假？")

            self.assertIn("试用期", response.answer)
            self.assertEqual(len(response.tool_calls), 2)
            self.assertEqual(response.tool_calls[0]["name"], "knowledge_retrieve")
            self.assertEqual(response.tool_calls[1]["name"], "export_report")
            self.assertIsNotNone(response.report_path)
            self.assertTrue(Path(response.report_path or "").is_file())
            self.assertEqual(len(agent.memory), 2)
            self.assertEqual(len(fake_llm.calls[0]["kwargs"]["tools"]), 3)

            second_call_messages = fake_llm.calls[1]["messages"]
            self.assertTrue(
                any(message.get("role") == "tool" for message in second_call_messages)
            )

    def test_evaluator_exports_summary_and_detail_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            settings = make_settings(root)
            settings.ensure_directories()
            dataset_path = root / "qa.csv"
            dataset_path.write_text(
                "question,expected_source,expected_answer\n"
                "员工何时申请年假？,employee_handbook.md,通过试用期后。\n",
                encoding="utf-8-sig",
            )
            evaluator = RetrievalEvaluator(
                settings,
                FakeVectorStore(),  # type: ignore[arg-type]
            )

            result = evaluator.evaluate(
                dataset_path,
                top_k=3,
                output_path=root / "eval.xlsx",
            )

            self.assertEqual(float(result.summary.iloc[0]["Hit@K"]), 1.0)
            self.assertIsNotNone(result.output_path)
            from openpyxl import load_workbook

            workbook = load_workbook(result.output_path or "")
            self.assertEqual(
                workbook.sheetnames,
                ["评测指标", "问题明细"],
            )
            self.assertEqual(workbook["问题明细"]["F2"].value, True)


if __name__ == "__main__":
    unittest.main()
