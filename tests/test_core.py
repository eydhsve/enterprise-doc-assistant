"""Core tests that do not require external model downloads."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from src.conversation_mem import ConversationMemory
from src.document_parser import DocumentParseError, DocumentParser, chunk_text
from src.evaluator import RetrievalEvaluator


class DocumentParserTests(unittest.TestCase):
    def test_chunk_text_has_overlap_and_multiple_chunks(self) -> None:
        text = "第一句。" * 160
        chunks = chunk_text(text, chunk_size=220, chunk_overlap=50)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(chunk.text for chunk in chunks))
        self.assertEqual(chunks[0].index, 0)
        self.assertLess(chunks[1].start_char, chunks[0].end_char)

    def test_parse_markdown_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.md"
            path.write_text("# 标题\n\n这是用于测试的文档内容。", encoding="utf-8")
            document = DocumentParser(chunk_size=100, chunk_overlap=20).parse_file(
                path
            )

            self.assertEqual(document.filename, "sample.md")
            self.assertIn("测试", document.text)
            self.assertGreaterEqual(len(document.chunks), 1)
            self.assertEqual(document.chunks[0].metadata, {})

    def test_missing_file_raises_clear_error(self) -> None:
        with self.assertRaises(DocumentParseError):
            DocumentParser().parse_file("missing-file.pdf")

    def test_parse_docx_paragraphs(self) -> None:
        try:
            from docx import Document as DocxDocument
        except ImportError:
            self.skipTest("python-docx is not installed.")

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "handbook.docx"
            document = DocxDocument()
            document.add_heading("员工手册", level=1)
            document.add_paragraph("员工通过试用期后可以申请年假。")
            document.save(path)

            parsed = DocumentParser().parse_file(path)

            self.assertIn("员工手册", parsed.text)
            self.assertIn("试用期", parsed.text)
            self.assertEqual(parsed.extension, ".docx")


class ConversationMemoryTests(unittest.TestCase):
    def test_history_is_trimmed_to_message_limit(self) -> None:
        memory = ConversationMemory(max_messages=4, max_context_chars=2000)
        memory.add_exchange("问题1", "回答1")
        memory.add_exchange("问题2", "回答2")
        memory.add_exchange("问题3", "回答3")

        snapshot = memory.snapshot()
        self.assertEqual(len(snapshot), 4)
        self.assertEqual(snapshot[0]["content"], "问题2")
        self.assertEqual(snapshot[-1]["content"], "回答3")

    def test_context_keeps_system_and_current_question(self) -> None:
        memory = ConversationMemory(max_messages=4, max_context_chars=2000)
        memory.add_exchange("旧问题", "旧回答")
        messages = memory.context_messages("系统提示", "当前问题")

        self.assertEqual(messages[0]["role"], "system")
        self.assertEqual(messages[-1]["content"], "当前问题")

    def test_long_current_question_is_bounded(self) -> None:
        memory = ConversationMemory(max_messages=4, max_context_chars=1000)
        messages = memory.context_messages("系统提示", "长问题" * 2000)

        total_chars = sum(len(message["content"]) for message in messages)
        self.assertLessEqual(total_chars, 1000)


class EvaluatorTests(unittest.TestCase):
    def test_source_matching_accepts_filename_and_path(self) -> None:
        self.assertTrue(
            RetrievalEvaluator.source_matches(
                "employee_handbook.md",
                "/data/docs/employee_handbook.md",
                "employee_handbook.md",
            )
        )
        self.assertFalse(
            RetrievalEvaluator.source_matches(
                "security_policy.md",
                "/data/docs/employee_handbook.md",
                "employee_handbook.md",
            )
        )


if __name__ == "__main__":
    unittest.main()
